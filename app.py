import io
import re
import streamlit as st
from docx import Document
from docx.shared import Cm, Pt
from docx.oxml.ns import qn
from docx.oxml.shared import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH

#格式設定
CHINESE_UPPER = ["", "壹", "貳", "叁", "肆", "伍", "陸", "柒", "捌", "玖", "拾"]
CHINESE_LOWER = ["", "一", "二", "三", "四", "五", "六", "七", "八", "九", "十"]
BASE_OFFSET = 0.75
FORMAT_CONFIG = {
    "L1": {"left": 1.129 + BASE_OFFSET, "first": -1.129, "size": 16, "font": "微軟正黑體"},
    "L2": {"left": 1.693 + BASE_OFFSET, "first": -1.129, "size": 14, "font": "標楷體"}, 
    "L3": {"left": 2.187 + BASE_OFFSET, "first": -1.129, "size": 14, "font": "標楷體"}, 
    "L4": {"left": 2.681 + BASE_OFFSET, "first": -1.129, "size": 14, "font": "標楷體"}, 
    "L5": {"left": 3.175 + BASE_OFFSET, "first": -1.129, "size": 14, "font": "標楷體"}, 
    "L6": {"left": 3.669 + BASE_OFFSET, "first": -1.129, "size": 14, "font": "標楷體"}, 
}
SCRIPTURE_GAP = 2.0
SCRIPTURE_SIZE = 12
SCRIPTURE_FONT = "標楷體"

#工具
def set_run_font(run, font_name, font_size):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = Pt(font_size)

def apply_paragraph_settings(p):
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    pPr = p._element.get_or_add_pPr()
    contextualSpacing = OxmlElement('w:contextualSpacing')
    contextualSpacing.set(qn('w:val'), '1')
    pPr.append(contextualSpacing)
    snapToGrid = OxmlElement('w:snapToGrid')
    snapToGrid.set(qn('w:val'), '0')
    pPr.append(snapToGrid)

def to_fullwidth_punct(text):
    half_punct = ",.?!;:()[]{}'<>"
    full_punct = "，。？！；：（）［］｛｝’《》"
    table = str.maketrans(half_punct, full_punct)
    return text.translate(table)

def get_label(level, count):
    if level == "L1": return CHINESE_UPPER[count]
    if level == "L2": return CHINESE_LOWER[count]
    if level == "L3": return str(count)
    if level == "L4": return chr(96 + count)
    if level == "L5": return f"({CHINESE_LOWER[count]})"
    if level == "L6": return f"({count})"
    return ""

def reset_sub_counters(counters, current_level):
    levels = ["L1", "L2", "L3", "L4", "L5", "L6"]
    try:
        start_index = levels.index(current_level) + 1
        for i in range(start_index, len(levels)):
            counters[levels[i]] = 0
    except ValueError:
        pass

#資料處理
def generate_docx(input_text_string):
    doc = Document()
    counters = {"L1": 0, "L2": 0, "L3": 0, "L4": 0, "L5": 0, "L6": 0}
    current_level = "L1" 
    for line in input_text_string.split('\n'):
        clean_line = line.strip()
        if not clean_line: continue
        if clean_line.startswith("S:") or clean_line.startswith("S："):
            line_type = "S"
            text = clean_line[2:] 
        else:
            parts = re.split(r'[:：]', clean_line, maxsplit=1)
            if len(parts) == 2:
                line_type = parts[0]
                text = parts[1]
            else:
                continue
        p = doc.add_paragraph()
        apply_paragraph_settings(p)
        if line_type in FORMAT_CONFIG:
            current_level = line_type
            counters[current_level] += 1
            reset_sub_counters(counters, current_level)
            label = get_label(current_level, counters[current_level])
            text = to_fullwidth_punct(text)
            config = FORMAT_CONFIG[current_level]
            p.paragraph_format.left_indent = Cm(config["left"])
            p.paragraph_format.first_line_indent = Cm(config["first"])
            run = p.add_run(f"{label}\t{text}")
            set_run_font(run, config["font"], config["size"])
        elif line_type == "S":
            parent_left = FORMAT_CONFIG[current_level]["left"]
            p.paragraph_format.left_indent = Cm(parent_left + SCRIPTURE_GAP)
            p.paragraph_format.first_line_indent = Cm(-SCRIPTURE_GAP)
            parts = re.split(r'[|｜]', text, maxsplit=1)
            if len(parts) == 2:
                book_name = to_fullwidth_punct(parts[0])
                scripture = to_fullwidth_punct(parts[1])
                run = p.add_run(f"{book_name}\t{scripture}")
                set_run_font(run, SCRIPTURE_FONT, SCRIPTURE_SIZE)
    virtual_file = io.BytesIO()
    doc.save(virtual_file)
    virtual_file.seek(0)
    return virtual_file

#Streamlit介面
def add_more_blocks():
    """按鈕回呼函式：一次新增 3 行區塊"""
    st.session_state.block_count += 3

def main():
    st.set_page_config(page_title="綱要排版工具", layout="wide")
    st.title("綱要排版工具")
    st.subheader("檔案設定")
    custom_name = st.text_input("要輸出的檔名(不用副檔名)：", value="ex : 國高特綱要")
    output_filename = f"{custom_name.strip()}.docx" if custom_name.strip() else "自動排版綱要.docx"
    st.markdown("---")
    input_mode = st.radio("輸入方式：", ["一段一段打", "純文字 (需用特定指令開頭)"], horizontal=True)
    final_input_text = ""
    if "一段一段打" in input_mode:
        st.markdown("**使用方式：** 選擇階層並輸入內容。切換至 -S- 會展開經文輸入框。")
        
        if "block_count" not in st.session_state:
            st.session_state.block_count = 5 
            
        blocks_data = []
        for i in range(st.session_state.block_count):
            c1, c2 = st.columns([1, 6])
            
            with c1:
                lvl = st.selectbox("階層", ["大點", "中點", "小點", "英文小點(a)", "中文小小點(一)", "數字小小點(1)", "S"], key=f"lvl_{i}", label_visibility="collapsed")
                
            with c2:
                if lvl == "S":
                    sc1, sc2 = st.columns([1, 3])
                    with sc1:
                        book = st.text_input("卷名", key=f"book_{i}", label_visibility="collapsed", placeholder="卷名 (例: 創2:9)")
                    with sc2:
                        content = st.text_input("經文", key=f"scripture_{i}", label_visibility="collapsed", placeholder="經文... 留白自動跳過")
                    
                    # 背景自動補上分隔符號 |
                    if book.strip() or content.strip():
                        txt = f"{book.strip()}|{content.strip()}"
                    else:
                        txt = ""
                else:
                    # 一般階層的單一輸入框
                    txt = st.text_input("內容", key=f"txt_{i}", label_visibility="collapsed", placeholder="輸入內容... 留白自動跳過")
                    
            blocks_data.append((lvl, txt))
        st.button("新增 3 行", on_click=add_more_blocks)

        LEVEL_MAP = {
            "大點": "L1",
            "中點": "L2",
            "小點": "L3",
            "英文小點(a)": "L4",
            "中文小小點(一)": "L5",
            "數字小小點(1)": "L6",
            "S": "S"
        }
        valid_lines = [f"{LEVEL_MAP.get(lvl, lvl)}:{txt.strip()}" for lvl, txt in blocks_data if txt.strip()]
        final_input_text = "\n".join(valid_lines)
        
    else:
        st.markdown("💡 **操作說明：** 直接貼上特殊指令開頭的純文字")
        final_input_text = st.text_area(
            "綱要：", 
            height=400, 
            placeholder="照以下格式輸入：\nL1:大點內容\nL2:中點內容\nL3:\n...~\nL6:\nS:創2:9|經文內容"
        )

    st.markdown("---")
    
    # --- 生成與下載狀態控制 ---
    if st.button("生成 Word 檔"):
        if not final_input_text.strip():
            st.error("錯誤：請先輸入內容！")
        else:
            try:
                # 將產生的檔案與檔名存入 session_state，避免網頁重整時消失
                st.session_state.generated_docx = generate_docx(final_input_text)
                st.session_state.generated_filename = output_filename
            except Exception as e:
                st.error(f"發生未預期的錯誤：{e}")

    # 若記憶體中有生成的檔案，則渲染下載按鈕
    if "generated_docx" in st.session_state:
        st.success("完成！")
        st.download_button(
            label="下載檔案",
            data=st.session_state.generated_docx,
            file_name=st.session_state.generated_filename,
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

if __name__ == "__main__":
    main()